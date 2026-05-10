# modules/excel_executor.py
import logging
import re

logger = logging.getLogger("OfficeAgent")


def _xl_color(hex_color):
    from openpyxl.styles import Color
    return Color(rgb=hex_color.lstrip("#").upper().zfill(6))


class ExcelExecutor:
    def __init__(self, wb, ws):
        self.wb = wb
        self.ws = ws

    def _iter_cells(self, range_ref):
        ref = (range_ref or "").strip()
        if not ref:
            return []
        if ":" not in ref:
            return [self.ws[ref]]
        cells = []
        for row in self.ws[ref]:
            for cell in row:
                cells.append(cell)
        return cells

    def run(self, action_dict):
        action  = action_dict.get("action")
        handler = getattr(self, f"_do_{action}", None)
        if not handler:
            logger.warning(f"Excel: Unknown action '{action}'")
            return False
        try:
            handler(action_dict)
            return True
        except Exception as e:
            logger.error(f"Excel action '{action}' failed: {e}")
            return False

    # â”€â”€ Cell / Range Operations â”€â”€â”€â”€â”€â”€â”€â”€â”€â”€â”€â”€â”€â”€â”€â”€â”€â”€â”€â”€â”€â”€â”€â”€â”€â”€â”€â”€â”€â”€â”€â”€â”€â”€â”€

    def _do_write_cell(self, p):
        self.ws[p["cell"]] = p.get("value", "")

    def _do_write_formula(self, p):
        self.ws[p["cell"]] = p.get("formula", "")

    def _do_write_range(self, p):
        import openpyxl
        start = p.get("start_cell", "A1")
        values = p.get("values", [])
        if isinstance(values, list):
            for i, row_data in enumerate(values):
                if isinstance(row_data, list):
                    for j, val in enumerate(row_data):
                        cell = self.ws.cell(
                            row=openpyxl.utils.cell.coordinate_to_tuple(start)[0] + i,
                            column=openpyxl.utils.cell.coordinate_to_tuple(start)[1] + j
                        )
                        cell.value = val
                else:
                    self.ws.cell(
                        row=openpyxl.utils.cell.coordinate_to_tuple(start)[0] + i,
                        column=openpyxl.utils.cell.coordinate_to_tuple(start)[1]
                    ).value = row_data

    def _do_read_cell(self, p):
        val = self.ws[p["cell"]].value
        logger.info(f"Cell {p['cell']} = {val}")
        return val

    def _do_clear_range(self, p):
        for row in self.ws[p["range"]]:
            for cell in row:
                cell.value = None

    def _do_clear_all(self, p):
        for row in self.ws.iter_rows():
            for cell in row:
                cell.value = None

    def _do_clear_format(self, p):
        from openpyxl.styles import Font, PatternFill, Alignment, Border
        for row in self.ws[p["range"]]:
            for cell in row:
                cell.font      = Font()
                cell.fill      = PatternFill()
                cell.alignment = Alignment()
                cell.border    = Border()

    def _do_copy_range(self, p):
        logger.info(f"Copy range {p.get('range')} (clipboard â€” requires win32com)")

    def _do_cut_range(self, p):
        logger.info(f"Cut range {p.get('range')} (clipboard â€” requires win32com)")

    def _do_paste_range(self, p):
        logger.info(f"Paste to {p.get('cell')} (clipboard â€” requires win32com)")

    def _do_paste_values_only(self, p):
        logger.info(f"Paste values only to {p.get('cell')} (requires win32com)")

    def _do_undo(self, p):
        logger.info("Undo (requires win32com)")

    def _do_redo(self, p):
        logger.info("Redo (requires win32com)")

    # â”€â”€ Font & Style â”€â”€â”€â”€â”€â”€â”€â”€â”€â”€â”€â”€â”€â”€â”€â”€â”€â”€â”€â”€â”€â”€â”€â”€â”€â”€â”€â”€â”€â”€â”€â”€â”€â”€â”€â”€â”€â”€â”€â”€â”€â”€â”€â”€â”€

    def _do_set_bold(self, p):
        from openpyxl.styles import Font
        for cell in self._iter_cells(p["range"]):
            cell.font = Font(
                bold=p.get("bold", True),
                italic=cell.font.italic,
                size=cell.font.size,
                name=cell.font.name,
                color=cell.font.color
            )

    def _do_set_italic(self, p):
        from openpyxl.styles import Font
        for cell in self._iter_cells(p["range"]):
            cell.font = Font(
                italic=p.get("italic", True),
                bold=cell.font.bold,
                size=cell.font.size,
                name=cell.font.name
            )

    def _do_set_underline(self, p):
        from openpyxl.styles import Font
        val = "single" if p.get("underline", True) else None
        for row in self.ws[p["range"]]:
            for cell in row:
                cell.font = Font(
                    underline=val,
                    bold=cell.font.bold,
                    italic=cell.font.italic,
                    size=cell.font.size
                )

    def _do_set_strikethrough(self, p):
        from openpyxl.styles import Font
        for row in self.ws[p["range"]]:
            for cell in row:
                cell.font = Font(
                    strike=True,
                    bold=cell.font.bold,
                    italic=cell.font.italic,
                    size=cell.font.size
                )

    def _do_set_font_size(self, p):
        from openpyxl.styles import Font
        for cell in self._iter_cells(p["range"]):
            cell.font = Font(
                size=int(p["size"]),
                bold=cell.font.bold,
                italic=cell.font.italic,
                name=cell.font.name
            )

    def _do_set_font_name(self, p):
        from openpyxl.styles import Font
        for cell in self._iter_cells(p["range"]):
            cell.font = Font(
                name=p["name"],
                bold=cell.font.bold,
                italic=cell.font.italic,
                size=cell.font.size
            )

    def _do_set_font_color(self, p):
        from openpyxl.styles import Font
        color = _xl_color(p["color"])
        for cell in self._iter_cells(p["range"]):
            cell.font = Font(
                color=color,
                bold=cell.font.bold,
                italic=cell.font.italic,
                size=cell.font.size,
                name=cell.font.name
            )

    def _do_set_bg_color(self, p):
        from openpyxl.styles import PatternFill
        fill = PatternFill(
            start_color=p["color"],
            end_color=p["color"],
            fill_type="solid"
        )
        for cell in self._iter_cells(p["range"]):
            cell.fill = fill

    def _do_set_border(self, p):
        from openpyxl.styles import Border, Side
        side   = Side(style=p.get("style", "thin"))
        border = Border(left=side, right=side, top=side, bottom=side)
        for row in self.ws[p["range"]]:
            for cell in row:
                cell.border = border

    def _do_remove_border(self, p):
        from openpyxl.styles import Border
        for row in self.ws[p["range"]]:
            for cell in row:
                cell.border = Border()

    def _do_set_alignment(self, p):
        from openpyxl.styles import Alignment
        horiz = p.get("alignment", "left")
        for row in self.ws[p["range"]]:
            for cell in row:
                cell.alignment = Alignment(
                    horizontal=horiz,
                    vertical=cell.alignment.vertical,
                    wrap_text=cell.alignment.wrap_text
                )

    def _do_set_vertical_alignment(self, p):
        from openpyxl.styles import Alignment
        for row in self.ws[p["range"]]:
            for cell in row:
                cell.alignment = Alignment(
                    vertical=p.get("alignment", "center"),
                    horizontal=cell.alignment.horizontal
                )

    def _do_set_wrap_text(self, p):
        from openpyxl.styles import Alignment
        for row in self.ws[p["range"]]:
            for cell in row:
                cell.alignment = Alignment(
                    wrap_text=p.get("wrap", True),
                    horizontal=cell.alignment.horizontal
                )

    def _do_set_number_format(self, p):
        for cell in self._iter_cells(p["range"]):
            cell.number_format = p.get("format", "General")

    # â”€â”€ Merge / Unmerge â”€â”€â”€â”€â”€â”€â”€â”€â”€â”€â”€â”€â”€â”€â”€â”€â”€â”€â”€â”€â”€â”€â”€â”€â”€â”€â”€â”€â”€â”€â”€â”€â”€â”€â”€â”€â”€â”€â”€â”€â”€â”€

    def _do_merge_cells(self, p):
        self.ws.merge_cells(p["range"])

    def _do_unmerge_cells(self, p):
        self.ws.unmerge_cells(p["range"])

    # â”€â”€ Rows / Columns â”€â”€â”€â”€â”€â”€â”€â”€â”€â”€â”€â”€â”€â”€â”€â”€â”€â”€â”€â”€â”€â”€â”€â”€â”€â”€â”€â”€â”€â”€â”€â”€â”€â”€â”€â”€â”€â”€â”€â”€â”€â”€â”€

    def _do_insert_row(self, p):
        self.ws.insert_rows(int(p["row"]))

    def _do_insert_column(self, p):
        self.ws.insert_cols(ord(p["column"].upper()) - 64)

    def _do_delete_row(self, p):
        self.ws.delete_rows(int(p["row"]))

    def _do_delete_column(self, p):
        self.ws.delete_cols(ord(p["column"].upper()) - 64)

    def _do_set_row_height(self, p):
        self.ws.row_dimensions[int(p["row"])].height = float(p["height"])

    def _do_set_column_width(self, p):
        self.ws.column_dimensions[p["column"].upper()].width = float(p["width"])

    def _do_autofit_columns(self, p):
        for col in self.ws.columns:
            max_len   = max((len(str(c.value or "")) for c in col), default=0)
            col_letter = col[0].column_letter
            self.ws.column_dimensions[col_letter].width = min(max_len + 4, 60)

    def _do_autofit_rows(self, p):
        logger.info("Autofit rows (approximate â€” full support requires win32com)")
        for row in self.ws.iter_rows():
            self.ws.row_dimensions[row[0].row].height = 15

    def _do_hide_row(self, p):
        self.ws.row_dimensions[int(p["row"])].hidden = True

    def _do_unhide_row(self, p):
        self.ws.row_dimensions[int(p["row"])].hidden = False

    def _do_hide_column(self, p):
        self.ws.column_dimensions[p["column"].upper()].hidden = True

    def _do_unhide_column(self, p):
        self.ws.column_dimensions[p["column"].upper()].hidden = False

    # â”€â”€ Sheet Operations â”€â”€â”€â”€â”€â”€â”€â”€â”€â”€â”€â”€â”€â”€â”€â”€â”€â”€â”€â”€â”€â”€â”€â”€â”€â”€â”€â”€â”€â”€â”€â”€â”€â”€â”€â”€â”€â”€â”€â”€â”€

    def _do_add_sheet(self, p):
        self.wb.create_sheet(title=p.get("name", "Sheet"))

    def _do_delete_sheet(self, p):
        name = p.get("name")
        if name in self.wb.sheetnames:
            del self.wb[name]

    def _do_rename_sheet(self, p):
        old_name = p.get("old_name")
        new_name = p.get("new_name", "Sheet")
        if old_name and old_name in self.wb.sheetnames:
            self.wb[old_name].title = new_name
            return
        # If old name is not provided, rename the active sheet.
        self.ws.title = new_name

    def _do_duplicate_sheet(self, p):
        src = self.wb[p.get("name", self.ws.title)]
        self.wb.copy_worksheet(src)

    def _do_hide_sheet(self, p):
        self.wb[p["name"]].sheet_state = "hidden"

    def _do_unhide_sheet(self, p):
        self.wb[p["name"]].sheet_state = "visible"

    def _do_set_active_sheet(self, p):
        self.wb.active = self.wb[p["name"]]

    def _do_move_sheet(self, p):
        self.wb.move_sheet(p.get("name"), offset=int(p.get("position", 0)))

    def _do_protect_sheet(self, p):
        self.ws.protection.sheet    = True
        self.ws.protection.password = p.get("password", "")

    def _do_unprotect_sheet(self, p):
        self.ws.protection.sheet    = False
        self.ws.protection.password = ""

    def _do_protect_workbook(self, p):
        logger.info("Workbook protection (requires win32com)")

    def _do_unprotect_workbook(self, p):
        logger.info("Workbook unprotection (requires win32com)")

    # â”€â”€ Data Tools â”€â”€â”€â”€â”€â”€â”€â”€â”€â”€â”€â”€â”€â”€â”€â”€â”€â”€â”€â”€â”€â”€â”€â”€â”€â”€â”€â”€â”€â”€â”€â”€â”€â”€â”€â”€â”€â”€â”€â”€â”€â”€â”€â”€â”€â”€â”€

    def _do_freeze_panes(self, p):
        self.ws.freeze_panes = p["cell"]

    def _do_unfreeze_panes(self, p):
        self.ws.freeze_panes = None

    def _do_sort_range(self, p):
        reverse = p.get("order", "ascending") == "descending"
        region  = list(self.ws[p["range"]])
        region.sort(key=lambda r: r[0].value or "", reverse=reverse)

    def _do_filter_range(self, p):
        self.ws.auto_filter.ref = p["range"]

    def _do_remove_filter(self, p):
        self.ws.auto_filter.ref = None

    def _do_find_replace(self, p):
        for row in self.ws.iter_rows():
            for cell in row:
                if cell.value and str(cell.value) == p.get("find_text", ""):
                    cell.value = p.get("replace_text", "")

    def _do_remove_duplicates(self, p):
        seen, rows_to_delete = set(), []
        for row in self.ws.iter_rows():
            key = tuple(c.value for c in row)
            if key in seen:
                rows_to_delete.append(row[0].row)
            else:
                seen.add(key)
        for row_idx in reversed(rows_to_delete):
            self.ws.delete_rows(row_idx)

    def _do_text_to_columns(self, p):
        logger.info(f"Text to columns with delimiter '{p.get('delimiter')}' (requires win32com)")

    def _do_create_named_range(self, p):
        self.wb.defined_names[p["name"]] = p["range"]

    def _do_add_conditional_formatting(self, p):
        from openpyxl.formatting.rule import ColorScaleRule
        rule = ColorScaleRule(
            start_type="min", start_color="FF0000",
            end_type="max",   end_color="00FF00"
        )
        self.ws.conditional_formatting.add(p["range"], rule)

    def _do_add_data_validation(self, p):
        from openpyxl.worksheet.datavalidation import DataValidation
        vals = p.get("values", [])
        if isinstance(vals, list):
            vals = ",".join(str(v) for v in vals)
        dv = DataValidation(type="list", formula1=f'"{vals}"')
        self.ws.add_data_validation(dv)
        dv.add(self.ws[p["range"]])

    def _do_insert_comment(self, p):
        from openpyxl.comments import Comment
        self.ws[p["cell"]].comment = Comment(p.get("text", ""), "Agent")

    def _do_delete_comment(self, p):
        self.ws[p["cell"]].comment = None

    def _do_insert_hyperlink(self, p):
        cell           = self.ws[p["cell"]]
        cell.value     = p.get("text", p["cell"])
        cell.hyperlink = p.get("url", "")

    def _do_create_table(self, p):
        from openpyxl.worksheet.table import Table, TableStyleInfo
        from openpyxl.utils.cell import coordinate_to_tuple

        start = p.get("start_cell", "A1")
        rows = max(2, int(p.get("rows", 5)))
        cols = max(1, int(p.get("cols", 3)))

        # Excel tables require a header row; generate defaults when missing.
        start_row, start_col = coordinate_to_tuple(start)
        for i in range(cols):
            header_cell = self.ws.cell(row=start_row, column=start_col + i)
            if header_cell.value is None or str(header_cell.value).strip() == "":
                header_cell.value = f"Column{i + 1}"

        end     = self._offset_cell(start, rows - 1, cols - 1)
        ref     = f"{start}:{end}"
        tbl     = Table(displayName=f"Table{len(self.ws.tables) + 1}", ref=ref)
        tbl.tableStyleInfo = TableStyleInfo(
            name="TableStyleMedium9", showRowStripes=True
        )
        self.ws.add_table(tbl)

    def _do_insert_chart(self, p):
        from openpyxl.chart import BarChart, LineChart, PieChart, Reference
        chart_map = {
            "bar":  BarChart,
            "line": LineChart,
            "pie":  PieChart,
        }
        ChartClass = chart_map.get(p.get("chart_type", "bar"), BarChart)
        chart      = ChartClass()
        chart.title = "Chart"
        self.ws.add_chart(chart, p.get("start_cell", "E1"))

    def _do_create_pivot_table(self, p):
        logger.info(f"Pivot table from {p.get('source_range')} (requires win32com)")

    def _do_group_rows(self, p):
        self.ws.row_dimensions.group(int(p["start_row"]), int(p["end_row"]), hidden=False)

    def _do_ungroup_rows(self, p):
        self.ws.row_dimensions.group(int(p["start_row"]), int(p["end_row"]), hidden=False)

    def _do_insert_image(self, p):
        from openpyxl.drawing.image import Image
        import os
        path = p.get("path", "")
        if path and os.path.exists(path):
            img = Image(path)
            self.ws.add_image(img, p.get("cell", "A1"))

    # â”€â”€ View / Print â”€â”€â”€â”€â”€â”€â”€â”€â”€â”€â”€â”€â”€â”€â”€â”€â”€â”€â”€â”€â”€â”€â”€â”€â”€â”€â”€â”€â”€â”€â”€â”€â”€â”€â”€â”€â”€â”€â”€â”€â”€â”€â”€â”€â”€

    def _do_set_zoom(self, p):
        self.ws.sheet_view.zoomScale = int(p.get("level", 100))

    def _do_set_print_area(self, p):
        self.ws.print_area = p["range"]

    def _do_set_print_setup(self, p):
        self.ws.page_setup.orientation = p.get("orientation", "portrait")

    def _do_spell_check(self, p):
        logger.info("Spell check (requires win32com)")

    # â”€â”€ Save / Workbook â”€â”€â”€â”€â”€â”€â”€â”€â”€â”€â”€â”€â”€â”€â”€â”€â”€â”€â”€â”€â”€â”€â”€â”€â”€â”€â”€â”€â”€â”€â”€â”€â”€â”€â”€â”€â”€â”€â”€â”€â”€â”€


    def _do_create_workbook(self, p):
        from openpyxl import Workbook
        self.wb = Workbook()
        self.ws = self.wb.active

    def _do_open_workbook(self, p):
        # Open/load is handled by server-level lifecycle.
        return

    def _do_save_workbook(self, p):
        path = getattr(self.wb, "_path", "output.xlsx")
        self.wb.save(path)

    def _do_save_workbook_as(self, p):
        filename = p.get("filename", "output")
        if not filename.endswith(".xlsx"):
            filename += ".xlsx"
        self.wb.save(filename)

    def _do_close_workbook(self, p):
        logger.info("Workbook closed (session ended)")

    # â”€â”€ Helpers â”€â”€â”€â”€â”€â”€â”€â”€â”€â”€â”€â”€â”€â”€â”€â”€â”€â”€â”€â”€â”€â”€â”€â”€â”€â”€â”€â”€â”€â”€â”€â”€â”€â”€â”€â”€â”€â”€â”€â”€â”€â”€â”€â”€â”€â”€â”€â”€â”€â”€

    def _offset_cell(self, cell_ref, row_offset, col_offset):
        m = re.match(r'([A-Z]+)(\d+)', cell_ref.upper())
        if not m:
            return cell_ref
        col     = m.group(1)
        row     = int(m.group(2))
        col_num = sum(
            (ord(c) - 64) * (26 ** i)
            for i, c in enumerate(reversed(col))
        )
        new_col_num = col_num + col_offset
        new_col     = ""
        while new_col_num > 0:
            new_col     = chr((new_col_num - 1) % 26 + 65) + new_col
            new_col_num = (new_col_num - 1) // 26
        return f"{new_col}{row + row_offset}"
